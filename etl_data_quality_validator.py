"""
ETL Data Quality Validator
==========================
Purpose : Source-to-target data quality validation for ETL pipelines.
          Covers row count reconciliation, null checks, duplicate detection,
          referential integrity, and value-range assertions.

Usage:
    python etl_data_quality_validator.py src.csv tgt.csv
    python etl_data_quality_validator.py src.csv tgt.csv --report dq_report.xlsx

Dependencies:
    pip install pandas openpyxl
"""

import argparse
import json
import pandas as pd
import logging
from datetime import datetime
from dataclasses import dataclass, field
from typing import Optional
from openpyxl.styles import Alignment, Font, PatternFill

# ─── Logging ──────────────────────────────────────────────────────────────────
logging.basicConfig(
    level=logging.INFO,
    format="%(asctime)s | %(levelname)-8s | %(message)s",
    datefmt="%Y-%m-%d %H:%M:%S",
)
log = logging.getLogger(__name__)


# ─── Result model ─────────────────────────────────────────────────────────────
@dataclass
class CheckResult:
    check_name: str
    status: str          # PASS / FAIL / WARN
    expected: str
    actual: str
    details: str = ""
    run_ts: str = field(default_factory=lambda: datetime.now().strftime("%Y-%m-%d %H:%M:%S"))


# ─── Core validator class ──────────────────────────────────────────────────────
class ETLDataQualityValidator:
    """
    Validates data quality between a source DataFrame and a target DataFrame.
    Mirrors checks performed during SIT/UAT in Informatica / Ab Initio pipelines.
    """

    def __init__(self, source: pd.DataFrame, target: pd.DataFrame, pipeline_name: str = "ETL_PIPELINE"):
        self.source = source
        self.target = target
        self.pipeline_name = pipeline_name
        self.results: list[CheckResult] = []

    # ── 1. Row Count Reconciliation ────────────────────────────────────────────
    def check_row_count(self) -> CheckResult:
        """Source and target must have the same number of rows."""
        src_count = len(self.source)
        tgt_count = len(self.target)
        status = "PASS" if src_count == tgt_count else "FAIL"
        result = CheckResult(
            check_name="Row Count Reconciliation",
            status=status,
            expected=str(src_count),
            actual=str(tgt_count),
            details=f"Delta: {abs(src_count - tgt_count)} rows",
        )
        self.results.append(result)
        log.info("[%s] Row Count — %s (src=%d, tgt=%d)", self.pipeline_name, status, src_count, tgt_count)
        return result

    # ── 2. Null / Missing Value Check ──────────────────────────────────────────
    def check_nulls(self, critical_columns: list[str]) -> list[CheckResult]:
        """Critical columns must have zero nulls in the target."""
        check_results = []
        for col in critical_columns:
            if col not in self.target.columns:
                r = CheckResult(
                    check_name=f"Null Check — {col}",
                    status="FAIL",
                    expected="0 nulls",
                    actual="Column missing",
                )
                self.results.append(r)
                check_results.append(r)
                continue
            null_count = self.target[col].isnull().sum()
            status = "PASS" if null_count == 0 else "FAIL"
            r = CheckResult(
                check_name=f"Null Check — {col}",
                status=status,
                expected="0 nulls",
                actual=f"{null_count} nulls",
            )
            self.results.append(r)
            check_results.append(r)
            log.info("[%s] Null check on '%s' — %s (%d nulls)", self.pipeline_name, col, status, null_count)
        return check_results

    # ── 3. Duplicate Key Detection ─────────────────────────────────────────────
    def check_duplicates(self, key_columns: list[str]) -> CheckResult:
        """Target must have no duplicate rows on the business key."""
        dup_count = self.target.duplicated(subset=key_columns).sum()
        status = "PASS" if dup_count == 0 else "FAIL"
        result = CheckResult(
            check_name=f"Duplicate Check — {key_columns}",
            status=status,
            expected="0 duplicates",
            actual=f"{dup_count} duplicates",
        )
        self.results.append(result)
        log.info("[%s] Duplicate check on %s — %s (%d dups)", self.pipeline_name, key_columns, status, dup_count)
        return result

    # ── 4. Column-Level Aggregation Reconciliation ────────────────────────────
    def check_aggregates(self, numeric_columns: list[str], tolerance_pct: float = 0.01) -> list[CheckResult]:
        """
        Source and target SUM must match within tolerance.
        Typical use: amount fields, balance totals after ETL transformation.
        """
        check_results = []
        for col in numeric_columns:
            if col not in self.source.columns or col not in self.target.columns:
                r = CheckResult(
                    check_name=f"Aggregate Check — {col}",
                    status="FAIL",
                    expected="Column present in both",
                    actual="Column missing in source or target",
                )
                self.results.append(r)
                check_results.append(r)
                continue
            src_sum = self.source[col].sum()
            tgt_sum = self.target[col].sum()
            delta_pct = abs(src_sum - tgt_sum) / (abs(src_sum) + 1e-9)
            status = "PASS" if delta_pct <= tolerance_pct else "FAIL"
            r = CheckResult(
                check_name=f"Aggregate Check — {col}",
                status=status,
                expected=f"SUM={src_sum:.2f} (±{tolerance_pct*100}%)",
                actual=f"SUM={tgt_sum:.2f} (delta={delta_pct*100:.3f}%)",
            )
            self.results.append(r)
            check_results.append(r)
            log.info("[%s] Aggregate '%s' — %s", self.pipeline_name, col, status)
        return check_results

    # ── 5. Value Range / Domain Validation ────────────────────────────────────
    def check_value_range(self, column: str, min_val=None, max_val=None,
                          allowed_values: Optional[set] = None) -> CheckResult:
        """Target column values must fall within expected domain."""
        if column not in self.target.columns:
            r = CheckResult(
                check_name=f"Value Range — {column}",
                status="FAIL",
                expected="Column present",
                actual="Column missing",
            )
            self.results.append(r)
            return r

        series = self.target[column].dropna()
        violations = 0

        if allowed_values:
            violations = (~series.isin(allowed_values)).sum()
            detail = f"Allowed: {allowed_values}"
        else:
            if min_val is not None:
                violations += (series < min_val).sum()
            if max_val is not None:
                violations += (series > max_val).sum()
            detail = f"Expected range: [{min_val}, {max_val}]"

        status = "PASS" if violations == 0 else "FAIL"
        r = CheckResult(
            check_name=f"Value Range — {column}",
            status=status,
            expected=detail,
            actual=f"{violations} violations",
        )
        self.results.append(r)
        log.info("[%s] Value range '%s' — %s (%d violations)", self.pipeline_name, column, status, violations)
        return r

    # ── 6. Referential Integrity ───────────────────────────────────────────────
    def check_referential_integrity(self, fk_column: str, ref_column: str) -> CheckResult:
        """All target keys must exist in the matching source-file column."""
        target_keys = set(self.target[fk_column].dropna().unique())
        ref_keys = set(self.source[ref_column].dropna().unique())
        orphans = target_keys - ref_keys
        status = "PASS" if not orphans else "FAIL"
        r = CheckResult(
            check_name=f"Referential Integrity — {fk_column} → {ref_column}",
            status=status,
            expected="No orphan keys",
            actual=f"{len(orphans)} orphan keys: {list(orphans)[:5]}{'...' if len(orphans) > 5 else ''}",
        )
        self.results.append(r)
        log.info("[%s] Ref integrity '%s' — %s", self.pipeline_name, fk_column, status)
        return r

    # ── Summary Report ─────────────────────────────────────────────────────────
    def summary(self) -> pd.DataFrame:
        """Return all check results as a DataFrame."""
        df = pd.DataFrame([vars(r) for r in self.results])
        total   = len(df)
        passed  = (df["status"] == "PASS").sum()
        failed  = (df["status"] == "FAIL").sum()
        warned  = (df["status"] == "WARN").sum()

        print("\n" + "=" * 70)
        print(f"  DATA QUALITY REPORT — {self.pipeline_name}")
        print(f"  Run: {datetime.now().strftime('%Y-%m-%d %H:%M:%S')}")
        print("=" * 70)
        print(f"  Total Checks : {total}")
        print(f"  PASSED       : {passed}  ✅")
        print(f"  FAILED       : {failed}  ❌")
        print(f"  WARNINGS     : {warned}  ⚠️")
        print("=" * 70)
        print(df[["check_name", "status", "expected", "actual", "details"]].to_string(index=False))
        print("=" * 70 + "\n")
        return df


def main(argv: Optional[list[str]] = None) -> None:
    parser = argparse.ArgumentParser(description="Compare source and target CSV files for ETL data quality.")
    parser.add_argument("source_file", help="Path to the source CSV file")
    parser.add_argument("target_file", help="Path to the target CSV file")
    parser.add_argument("--report", default="dq_report.xlsx", help="Path for the generated Excel report")
    args = parser.parse_args(argv)

    log.info("Starting ETL Data Quality Validation...")

    source_df = pd.read_csv(args.source_file)
    target_df = pd.read_csv(args.target_file)

    validator = ETLDataQualityValidator(
        source=source_df,
        target=target_df,
        pipeline_name="ACCOUNTS_DWH_LOAD",
    )

    # Run all checks
    validator.check_row_count()
    validator.check_nulls(critical_columns=["account_id", "balance", "currency"])
    validator.check_duplicates(key_columns=["account_id"])
    validator.check_aggregates(numeric_columns=["balance"], tolerance_pct=0.001)
    validator.check_value_range(
        column="currency",
        allowed_values={"USD", "GBP", "EUR", "SGD", "AED"}
    )
    validator.check_value_range(
        column="balance",
        min_val=0,
        max_val=10_000_000
    )
    validator.check_referential_integrity(
        fk_column="customer_id",
        ref_column="customer_id",
    )

    validator.summary()
    file_checks = build_file_checks(source_df, target_df)
    differences = build_differences(source_df, target_df, key_column="account_id")
    write_excel_report(file_checks, differences, args.report)
    log.info("Report saved to %s", args.report)


def build_file_checks(source: pd.DataFrame, target: pd.DataFrame) -> pd.DataFrame:
    allowed_currencies = {"USD", "GBP", "EUR", "SGD", "AED"}
    checks = []

    def add_check(name: str, field: str, src_status: str, src_details: str,
                  tgt_status: str, tgt_details: str) -> None:
        checks.append({
            "Check": name,
            "Field": field,
            "SRC Status": src_status,
            "SRC Details": src_details,
            "TGT Status": tgt_status,
            "TGT Details": tgt_details,
        })

    add_check("Row count", "All rows", "INFO", f"{len(source)} rows",
              "INFO", f"{len(target)} rows")

    for column in ("account_id", "customer_id", "balance", "currency"):
        results = []
        for frame in (source, target):
            if column not in frame.columns:
                results.append(("FAIL", "Column missing"))
            else:
                null_count = int(frame[column].isna().sum())
                results.append(("PASS" if null_count == 0 else "FAIL", f"{null_count} nulls"))
        add_check("Null check", column, results[0][0], results[0][1], results[1][0], results[1][1])

    results = []
    for frame in (source, target):
        if "account_id" not in frame.columns:
            results.append(("FAIL", "Column missing"))
        else:
            duplicate_count = int(frame.duplicated(subset=["account_id"]).sum())
            results.append(("PASS" if duplicate_count == 0 else "FAIL", f"{duplicate_count} duplicates"))
    add_check("Duplicate key check", "account_id", results[0][0], results[0][1], results[1][0], results[1][1])

    results = []
    for frame in (source, target):
        if "currency" not in frame.columns:
            results.append(("FAIL", "Column missing"))
        else:
            invalid = frame.loc[frame["currency"].notna() & ~frame["currency"].isin(allowed_currencies), "currency"]
            details = f"{len(invalid)} invalid values"
            if not invalid.empty:
                details += f": {invalid.astype(str).drop_duplicates().tolist()[:5]}"
            results.append(("PASS" if invalid.empty else "FAIL", details))
    add_check("Allowed values", "currency", results[0][0], results[0][1], results[1][0], results[1][1])

    results = []
    for frame in (source, target):
        if "balance" not in frame.columns:
            results.append(("FAIL", "Column missing"))
        else:
            values = pd.to_numeric(frame["balance"], errors="coerce")
            non_numeric = int((frame["balance"].notna() & values.isna()).sum())
            violations = int(((values < 0) | (values > 10_000_000)).sum())
            status = "PASS" if non_numeric == 0 and violations == 0 else "FAIL"
            results.append((status, f"{violations} out of range; {non_numeric} non-numeric"))
    add_check("Value range", "balance [0, 10000000]", results[0][0], results[0][1], results[1][0], results[1][1])

    return pd.DataFrame(checks)


def build_differences(source: pd.DataFrame, target: pd.DataFrame, key_column: str) -> pd.DataFrame:
    difference_columns = ["Key", "Issue", "Field", "SRC Value", "TGT Value"]
    differences = []

    for column in sorted(set(source.columns) - set(target.columns)):
        differences.append({"Key": "", "Issue": "Column missing in TGT", "Field": column,
                            "SRC Value": "Present", "TGT Value": ""})
    for column in sorted(set(target.columns) - set(source.columns)):
        differences.append({"Key": "", "Issue": "Column missing in SRC", "Field": column,
                            "SRC Value": "", "TGT Value": "Present"})

    if key_column not in source.columns or key_column not in target.columns:
        differences.append({
            "Key": "",
            "Issue": "Cannot compare rows",
            "Field": key_column,
            "SRC Value": "Present" if key_column in source.columns else "Missing",
            "TGT Value": "Present" if key_column in target.columns else "Missing",
        })
        return pd.DataFrame(differences, columns=difference_columns)

    source_rows = source.copy()
    target_rows = target.copy()
    source_rows["__occurrence"] = source_rows.groupby(key_column, dropna=False, sort=False).cumcount()
    target_rows["__occurrence"] = target_rows.groupby(key_column, dropna=False, sort=False).cumcount()
    source_columns = [column for column in source.columns if column != key_column]
    target_columns = [column for column in target.columns if column != key_column]
    source_rows = source_rows.rename(columns={column: f"__src_{column}" for column in source_columns})
    target_rows = target_rows.rename(columns={column: f"__tgt_{column}" for column in target_columns})
    merged = source_rows.merge(
        target_rows,
        on=[key_column, "__occurrence"],
        how="outer",
        indicator=True,
        sort=False,
    )

    for _, row in merged.iterrows():
        key = row[key_column]
        occurrence = int(row["__occurrence"]) + 1
        if row["_merge"] != "both":
            in_source = row["_merge"] == "left_only"
            row_values = {
                column: row[f"__src_{column}" if in_source else f"__tgt_{column}"]
                for column in (source_columns if in_source else target_columns)
            }
            differences.append({
                "Key": key,
                "Issue": "Row missing in TGT" if in_source else "Row missing in SRC",
                "Field": "Entire row",
                "SRC Value": json.dumps(row_values, default=str) if in_source else "",
                "TGT Value": "" if in_source else json.dumps(row_values, default=str),
            })
            continue

        for column in (column for column in source_columns if column in target_columns):
            src_value = row[f"__src_{column}"]
            tgt_value = row[f"__tgt_{column}"]
            both_null = pd.isna(src_value) and pd.isna(tgt_value)
            equal = both_null or (not pd.isna(src_value) and not pd.isna(tgt_value) and src_value == tgt_value)
            if not equal:
                differences.append({
                    "Key": key,
                    "Issue": "Value mismatch",
                    "Field": column,
                    "SRC Value": src_value,
                    "TGT Value": tgt_value,
                })

    if not differences:
        differences.append({"Key": "", "Issue": "No differences found", "Field": "", "SRC Value": "", "TGT Value": ""})
    return pd.DataFrame(differences, columns=difference_columns)


def write_excel_report(file_checks: pd.DataFrame, differences: pd.DataFrame, report_path: str) -> None:
    status_fills = {
        "PASS": PatternFill("solid", fgColor="E2F0D9"),
        "FAIL": PatternFill("solid", fgColor="FCE4D6"),
        "INFO": PatternFill("solid", fgColor="DDEBF7"),
    }

    with pd.ExcelWriter(report_path, engine="openpyxl") as writer:
        file_checks.to_excel(writer, sheet_name="File Checks", index=False)
        differences.to_excel(writer, sheet_name="Differences", index=False)

        for worksheet in writer.book.worksheets:
            worksheet.freeze_panes = "A2"
            worksheet.auto_filter.ref = worksheet.dimensions
            for cell in worksheet[1]:
                cell.font = Font(bold=True, color="FFFFFF")
                cell.fill = PatternFill("solid", fgColor="244062")
                cell.alignment = Alignment(wrap_text=True, vertical="top")
            for column_cells in worksheet.columns:
                width = min(max(max(len(str(cell.value or "")) for cell in column_cells) + 2, 12), 48)
                worksheet.column_dimensions[column_cells[0].column_letter].width = width
            for row in worksheet.iter_rows(min_row=2):
                for cell in row:
                    cell.alignment = Alignment(vertical="top", wrap_text=True)
                    if cell.value in status_fills:
                        cell.fill = status_fills[cell.value]


if __name__ == "__main__":
    main()
